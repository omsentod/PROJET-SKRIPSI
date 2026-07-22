import os
import math
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ==============================================================================
# 1. PARAMETER & KONFIGURASI GLOBAL
# ==============================================================================
RATIO_SCHEMES = {
    "A": (0.5, 1.5),
    "B": (0.7, 1.3),
    "C": (0.6, 1.4), # Skema utama skripsi
    "D": (0.5, 1.5),
}

CLUSTER_LABELS = {
    0: "Hemat",
    1: "Premium",
}

# ==============================================================================
# 2. ALGORITMA FUZZY C-MEANS (FCM) MANUAL (NUMPY MURNI)
# ==============================================================================
def fuzzy_c_means_manual(data, n_clusters=2, m=2.0, error=1e-5, max_iter=300, init_centroids=None, seed=42):
    n_samples = len(data)
    if len(data.shape) == 1:
        data = data.reshape(-1, 1)
    
    n_features = data.shape[1]
    np.random.seed(seed)
    
    centers_history = []
    
    if init_centroids is not None:
        init_centroids = np.array(init_centroids).reshape(n_clusters, n_features)
        init_sorted = np.sort(init_centroids.flatten())
        centers_history.append(init_sorted.copy())
        
        distances = np.zeros((n_clusters, n_samples))
        for j in range(n_clusters):
            distances[j, :] = np.sum((data - init_centroids[j, :])**2, axis=1)
        distances = np.fmax(distances, 1e-10)
        temp = 1.0 / distances
        U = temp / np.sum(temp, axis=0, keepdims=True)
    else:
        U = np.random.dirichlet(np.ones(n_clusters), size=n_samples).T
    
    centers = np.zeros((n_clusters, n_features))
    
    for iteration in range(max_iter):
        U_prev = U.copy()
        
        for j in range(n_clusters):
            numerator = np.sum((U[j, :] ** m)[:, np.newaxis] * data, axis=0)
            denominator = np.sum(U[j, :] ** m)
            centers[j, :] = numerator / np.fmax(denominator, 1e-10)
            
        # Catat centers di setiap iterasi (diurutkan agar selaras)
        sorted_centers = np.sort(centers.flatten())
        centers_history.append(sorted_centers.copy())
            
        dist_sq = np.zeros((n_clusters, n_samples))
        for j in range(n_clusters):
            dist_sq[j, :] = np.sum((data - centers[j, :])**2, axis=1)
            
        dist_sq = np.fmax(dist_sq, 1e-10)
        
        temp = dist_sq ** (-1.0 / (m - 1.0))
        U = temp / np.fmax(np.sum(temp, axis=0, keepdims=True), 1e-10)
        
        if np.linalg.norm(U - U_prev) < error:
            break
            
    labels = np.argmax(U, axis=0)
    return centers, U, labels, iteration + 1, centers_history

# Hitung Metrik Xie-Beni Manual untuk Laporan Metadata
def calculate_xie_beni_metrics(data, centers, U, m=2.0):
    n_samples = len(data)
    n_clusters = len(centers)
    if len(data.shape) == 1:
        data = data.reshape(-1, 1)
    if len(centers.shape) == 1:
        centers = centers.reshape(-1, 1)
        
    sigma = 0.0
    for j in range(n_clusters):
        d_sq = np.sum((data - centers[j, :])**2, axis=1)
        sigma += np.sum((U[j, :] ** m) * d_sq)
        
    sep = float('inf')
    for j in range(n_clusters):
        for k in range(j + 1, n_clusters):
            d_sq = np.sum((centers[j, :] - centers[k, :])**2)
            if d_sq < sep:
                sep = d_sq
                
    xb = sigma / (n_samples * sep) if sep != 0 and n_samples != 0 else float('inf')
    return xb, sigma, sep

def run_percentile_fcm(data_prices, m=2.0):
    q_vals = np.linspace(100 / 3, 200 / 3, 2)
    init_centers = np.percentile(data_prices, q_vals).reshape(-1, 1)
    
    centers, U, labels, iters, centers_history = fuzzy_c_means_manual(
        data_prices, n_clusters=2, m=m, init_centroids=init_centers
    )
    
    sorted_idx = np.argsort(centers.flatten())
    sorted_cntr = centers.flatten()[sorted_idx]
    sorted_u = U[sorted_idx]
    sorted_labels = np.argmax(sorted_u, axis=0)
    
    return sorted_cntr, sorted_u, sorted_labels, iters, centers_history

def run_budget_anchored_fcm(data_prices, budget, ratio_scheme="C", m=2.0):
    ratios = RATIO_SCHEMES[ratio_scheme]
    init_centers = np.array([budget * r for r in ratios]).reshape(-1, 1)
    
    centers, U, labels, iters, centers_history = fuzzy_c_means_manual(
        data_prices, n_clusters=2, m=m, init_centroids=init_centers
    )
    
    sorted_idx = np.argsort(centers.flatten())
    sorted_cntr = centers.flatten()[sorted_idx]
    sorted_u = U[sorted_idx]
    sorted_labels = np.argmax(sorted_u, axis=0)
    
    return sorted_cntr, sorted_u, sorted_labels, iters, centers_history

# Haversine
def haversine_distance(lat1, lon1, lat2, lon2):
    R = 6371.0
    lat1_r = math.radians(lat1)
    lat2_r = math.radians(lat2)
    dlat = math.radians(lat2 - lat1)
    dlon = math.radians(lon2 - lon1)
    a = (math.sin(dlat / 2) ** 2 +
         math.cos(lat1_r) * math.cos(lat2_r) * math.sin(dlon / 2) ** 2)
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
    return R * c * 1.3

def get_transport_cost(num_persons, distance_km):
    if num_persons <= 1:
        rate = 2250
        desc = "GoRide (1 orang)"
    elif num_persons <= 4:
        rate = 5150
        desc = "GoCar Standard (2-4 orang)"
    else:
        rate = 6000
        desc = "GoCar XL (5-6 orang)"
    return round(distance_km * rate), desc

def get_candidates(df, cluster_labels, centroids, anchor=None, ratios=None):
    candidates = {0: [], 1: []}
    for i in range(2):
        items_c = df[cluster_labels == i].copy()
        target_price = anchor * ratios[i] if anchor is not None else centroids[i]
        
        if items_c.empty:
            df_temp = df.copy()
            df_temp["dist_to_target"] = (df_temp["Estimasi_Harga"] - target_price).abs()
            best = df_temp.nsmallest(15, "dist_to_target")
        else:
            items_c["dist_to_target"] = (items_c["Estimasi_Harga"] - target_price).abs()
            best = items_c.nsmallest(15, "dist_to_target")
            
        candidates[i] = best.to_dict("records")
    return candidates

def get_budget_first_packages(df_wis, df_hot, df_kul, w_lbl, h_lbl, k_lbl, w_cntr, h_cntr, k_cntr, anchor_w, anchor_h, anchor_k, ratios):
    w_cand = get_candidates(df_wis, w_lbl, w_cntr, anchor_w, ratios)
    h_cand = get_candidates(df_hot, h_lbl, h_cntr, anchor_h, ratios)
    k_cand = get_candidates(df_kul, k_lbl, k_cntr, anchor_k, ratios)
    
    pkg_res = {0: [], 1: []}
    rooms = math.ceil(DEFAULT_PERSONS / 2.0)
    nights = max(DEFAULT_DURATION - 1, 1)
    
    for i in range(2):
        comb_list = []
        for h in h_cand[i]:
            for w in w_cand[i]:
                for k in k_cand[i]:
                    c_hotel = h["Estimasi_Harga"] * nights * rooms
                    c_wisata = w["Estimasi_Harga"] * DEFAULT_PERSONS
                    c_kuliner = k["Estimasi_Harga"] * DEFAULT_PERSONS * 3 * DEFAULT_DURATION
                    dist = haversine_distance(h["Latitude"], h["Longitude"], w["Latitude"], w["Longitude"]) + \
                           haversine_distance(w["Latitude"], w["Longitude"], k["Latitude"], k["Longitude"]) + \
                           haversine_distance(k["Latitude"], k["Longitude"], h["Latitude"], h["Longitude"])
                    c_trans, _ = get_transport_cost(DEFAULT_PERSONS, dist)
                    total = c_hotel + c_wisata + c_kuliner + c_trans
                    
                    if total <= DEFAULT_BUDGET:
                       comb_list.append((h, w, k, dist, total))
                       
        if i == 0:
            comb_list = sorted(comb_list, key=lambda x: x[3])
        else:
            comb_list = sorted(comb_list, key=lambda x: (-x[1].get("Rating", 0), -x[0].get("Estimasi_Harga", 0), x[3]))
            
        pkg_res[i] = comb_list[:5]
    return pkg_res

def get_flexible_packages(df_wis, df_hot, df_kul, w_lbl, h_lbl, k_lbl, w_cntr, h_cntr, k_cntr):
    w_cand = get_candidates(df_wis, w_lbl, w_cntr)
    h_cand = get_candidates(df_hot, h_lbl, h_cntr)
    k_cand = get_candidates(df_kul, k_lbl, k_cntr)
    
    pkg_res = {0: [], 1: []}
    rooms = math.ceil(DEFAULT_PERSONS / 2.0)
    nights = max(DEFAULT_DURATION - 1, 1)
    
    for i in range(2):
        comb_list = []
        for h in h_cand[i][:8]:
            for w in w_cand[i][:8]:
                for k in k_cand[i][:8]:
                    c_hotel = h["Estimasi_Harga"] * nights * rooms
                    c_wisata = w["Estimasi_Harga"] * DEFAULT_PERSONS
                    c_kuliner = k["Estimasi_Harga"] * DEFAULT_PERSONS * 3 * DEFAULT_DURATION
                    dist = haversine_distance(h["Latitude"], h["Longitude"], w["Latitude"], w["Longitude"]) + \
                           haversine_distance(w["Latitude"], w["Longitude"], k["Latitude"], k["Longitude"]) + \
                           haversine_distance(k["Latitude"], k["Longitude"], h["Latitude"], h["Longitude"])
                    c_trans, _ = get_transport_cost(DEFAULT_PERSONS, dist)
                    total = c_hotel + c_wisata + c_kuliner + c_trans
                    comb_list.append((h, w, k, dist, total))
                    
        if i == 0:
            comb_list = sorted(comb_list, key=lambda x: x[3])
        else:
            comb_list = sorted(comb_list, key=lambda x: (-x[1].get("Rating", 0), -x[0].get("Estimasi_Harga", 0), x[3]))
            
        pkg_res[i] = comb_list[:5]
    return pkg_res

def get_dest_first_packages(df_wis, df_hot, df_kul, h_lbl, k_lbl, h_cntr, k_cntr, locked_dest, anchor_h_dest, anchor_k_dest, ratios):
    h_cand = get_candidates(df_hot, h_lbl, h_cntr, anchor_h_dest, ratios)
    k_cand = get_candidates(df_kul, k_lbl, k_cntr, anchor_k_dest, ratios)
    
    pkg_res = {0: [], 1: []}
    rooms = math.ceil(DEFAULT_PERSONS / 2.0)
    nights = max(DEFAULT_DURATION - 1, 1)
    
    for i in range(2):
        comb_list = []
        for h in h_cand[i]:
            for k in k_cand[i]:
                c_hotel = h["Estimasi_Harga"] * nights * rooms
                c_wisata = locked_dest["Estimasi_Harga"] * DEFAULT_PERSONS
                c_kuliner = k["Estimasi_Harga"] * DEFAULT_PERSONS * 3 * DEFAULT_DURATION
                dist = haversine_distance(h["Latitude"], h["Longitude"], locked_dest["Latitude"], locked_dest["Longitude"]) + \
                       haversine_distance(locked_dest["Latitude"], locked_dest["Longitude"], k["Latitude"], k["Longitude"]) + \
                       haversine_distance(k["Latitude"], k["Longitude"], h["Latitude"], h["Longitude"])
                c_trans, _ = get_transport_cost(DEFAULT_PERSONS, dist)
                total = c_hotel + c_wisata + c_kuliner + c_trans
                
                if total <= DEFAULT_BUDGET:
                    comb_list.append((h, locked_dest, k, dist, total))
                    
        if i == 0:
            comb_list = sorted(comb_list, key=lambda x: x[3])
        else:
            comb_list = sorted(comb_list, key=lambda x: (-x[0].get("Estimasi_Harga", 0), x[3]))
            
        pkg_res[i] = comb_list[:5]
    return pkg_res

# ==============================================================================
# 3. LOAD DATASET
# ==============================================================================
print("📂 Memuat dataset bersih...")
df_wisata = pd.read_excel("wisata_clean.xlsx")
df_hotel = pd.read_excel("hotel_clean.xlsx")
df_kuliner = pd.read_excel("tempat_makan_clean.xlsx")

df_wisata["Estimasi_Harga"] = df_wisata["Estimasi_Harga"].astype(float)
df_hotel["Estimasi_Harga"] = df_hotel["Estimasi_Harga"].astype(float)
df_kuliner["Estimasi_Harga"] = df_kuliner["Estimasi_Harga"].astype(float)

# ==============================================================================
# PARAMETER GLOBAL & PRE-COMPUTING FCM (DEFAULT VALUES FOR PIPELINE)
# ==============================================================================
DEFAULT_BUDGET = 1500000.0
DEFAULT_PERSONS = 1   # Skenario 1 orang
DEFAULT_DURATION = 2  # Skenario 2 hari (1 Malam)
DEFAULT_SCHEME = "C"

# ID Destinasi Ekstrem Wisata
CHEAPEST_WISATA_ID = 40     # Gua Maria Sendang Purwaningsih, HTM Rp 0
EXPENSIVE_WISATA_ID = 159   # Pujon Adventure & Rafting, HTM Rp 275.000

# Folder Output
output_dir = "ranking dataset"
if not os.path.exists(output_dir):
    os.makedirs(output_dir)
    print(f"📁 Folder '{output_dir}' berhasil dibuat.")

# Styles (Format ARGB 8-Karakter FF...)
font_title = Font(name="Segoe UI", size=14, bold=True, color="FFFFFFFF")
font_section = Font(name="Segoe UI", size=11, bold=True, color="FF1F4E79")
font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFFFF")
font_bold = Font(name="Segoe UI", size=10, bold=True)
font_regular = Font(name="Segoe UI", size=10)

fill_dark = PatternFill(start_color="FF2E4053", end_color="FF2E4053", fill_type="solid")
fill_light_blue = PatternFill(start_color="FFD6EAF8", end_color="FFD6EAF8", fill_type="solid")
fill_light_teal = PatternFill(start_color="FFD1F2EB", end_color="FFD1F2EB", fill_type="solid")
fill_soft_gray = PatternFill(start_color="FFF2F4F4", end_color="FFF2F4F4", fill_type="solid")

fill_hemat = PatternFill(start_color="FFD4EFDF", end_color="FFD4EFDF", fill_type="solid")
fill_balanced = PatternFill(start_color="FFFCF3CF", end_color="FFFCF3CF", fill_type="solid")
fill_premium = PatternFill(start_color="FFFADBD8", end_color="FFFADBD8", fill_type="solid")

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_right = Alignment(horizontal="right", vertical="center")

thin_border = Border(
    left=Side(style='thin', color='FFBDC3C7'),
    right=Side(style='thin', color='FFBDC3C7'),
    top=Side(style='thin', color='FFBDC3C7'),
    bottom=Side(style='thin', color='FFBDC3C7')
)

def auto_fit_columns(ws, min_width=10, max_width=45):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.coordinate in ws.merged_cells:
                continue
            val_str = str(cell.value or '')
            if val_str.startswith('='):
                val_str = "Formula_Cell" 
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, min_width), max_width)

# ==============================================================================
# CORE GENERATION FUNCTION
# ==============================================================================
def generate_ranking_excel(filename, workflow_title, wisata_data, hotel_data, kuliner_data, pkg_res):
    wb = openpyxl.Workbook()
    # Hapus sheet default pertama
    wb.remove(wb.active)
    
    datasets_info = [
        ("Wisata_Ranking", wisata_data, False, False),
        ("Hotel_Ranking", hotel_data, True, False),
        ("Kuliner_Ranking", kuliner_data, False, True)
    ]
    
    for sheet_name, data_pkg, is_hotel, is_kuliner in datasets_info:
        ws = wb.create_sheet(sheet_name)
        ws.views.sheetView[0].showGridLines = True
        
        # Banner Atas
        ws.merge_cells("A1:J2")
        ws["A1"] = f"HASIL PERANGKINGAN DATASET - {sheet_name.replace('_', ' ').upper()}"
        ws["A1"].font = font_title
        ws["A1"].fill = fill_dark
        ws["A1"].alignment = align_center
        
        # Metadata / Parameter Box
        ws["A4"] = "METADATA & PARAMETER ALGORITMA FCM"
        ws["A4"].font = font_section
        ws.merge_cells("A4:C4")
        
        # Ambil metadata riel dari data_pkg
        n_val = data_pkg["n"]
        iters_val = data_pkg["iters"]
        sigma_val = data_pkg["sigma"]
        sep_val = data_pkg["sep"]
        xb_val = data_pkg["xb"]
        
        metadata = [
            ("Banyak Sampel (n)", n_val, "C5"),
            ("Jumlah Iterasi FCM (t)", iters_val, "C6"),
            ("Total Variansi (Sigma) [Eq. 2.6]", sigma_val, "C7"),
            ("Separasi Centroid (sep) [Eq. 2.7]", sep_val, "C8"),
            ("Xie-Beni Index (XB) [Eq. 2.8]", xb_val, "C9")
        ]
        
        for label, val, cell_coord in metadata:
            r = int(cell_coord[1])
            ws.cell(row=r, column=1, value=label).font = font_bold
            ws.cell(row=r, column=1).border = thin_border
            cell_v = ws.cell(row=r, column=3, value=val)
            cell_v.font = font_bold
            cell_v.alignment = align_center
            cell_v.border = thin_border
            if label.startswith("Total") or label.startswith("Separasi"):
                cell_v.number_format = "#,##0.0"
                cell_v.alignment = align_right
            elif label.startswith("Xie-Beni"):
                cell_v.number_format = "0.000000"
                cell_v.alignment = align_right
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
            
        # Centroid Table
        ws["E4"] = "PUSAT KLASTER (CENTROID FINAL)"
        ws["E4"].font = font_section
        ws.merge_cells("E4:G4")
        
        ws.cell(row=5, column=5, value="Kelas Paket").font = font_header
        ws.cell(row=5, column=5).fill = fill_dark
        ws.cell(row=5, column=5).alignment = align_center
        ws.cell(row=5, column=5).border = thin_border
        
        ws.cell(row=5, column=6, value="Centroid Final (Rp)").font = font_header
        ws.cell(row=5, column=6).fill = fill_dark
        ws.cell(row=5, column=6).alignment = align_center
        ws.cell(row=5, column=6).border = thin_border
        
        ws.cell(row=5, column=7, value="Target Anchor (Rp)").font = font_header
        ws.cell(row=5, column=7).fill = fill_dark
        ws.cell(row=5, column=7).alignment = align_center
        ws.cell(row=5, column=7).border = thin_border
        
        for idx in range(2):
            r = 6 + idx
            lbl = CLUSTER_LABELS[idx].upper()
            cell_lbl = ws.cell(row=r, column=5, value=lbl)
            cell_lbl.font = font_bold
            cell_lbl.border = thin_border
            if idx == 0: cell_lbl.fill = fill_hemat
            else: cell_lbl.fill = fill_premium
            
            cell_c = ws.cell(row=r, column=6, value=float(data_pkg["centroids"][idx]))
            cell_c.font = font_bold
            cell_c.number_format = "Rp #,##0"
            cell_c.alignment = align_right
            cell_c.border = thin_border
            
            cell_a = ws.cell(row=r, column=7, value=float(data_pkg["anchors"][idx]) if data_pkg["anchors"] else "-")
            cell_a.font = font_regular
            if data_pkg["anchors"]:
                cell_a.number_format = "Rp #,##0"
            cell_a.alignment = align_right
            cell_a.border = thin_border
            
        # Riwayat Iterasi Centroid Table (Kolom M s.d. O)
        ws["M4"] = "RIWAYAT ITERASI PUSAT KLASTER (CENTROID HISTORY)"
        ws["M4"].font = font_section
        ws.merge_cells("M4:O4")
        
        hist_headers = ["Iterasi", "Centroid Hemat (Rp)", "Centroid Premium (Rp)"]
        for c_idx, h in enumerate(hist_headers):
            cell_h = ws.cell(row=5, column=13+c_idx, value=h)
            cell_h.font = font_header
            cell_h.fill = fill_dark
            cell_h.alignment = align_center
            cell_h.border = thin_border
            
        centers_history = data_pkg.get("centers_history", [])
        for iter_idx, centers_step in enumerate(centers_history):
            r = 6 + iter_idx
            
            # Kolom M (Iterasi)
            label_iter = f"{iter_idx} (Awal)" if iter_idx == 0 else str(iter_idx)
            cell_it = ws.cell(row=r, column=13, value=label_iter)
            cell_it.font = font_bold
            cell_it.alignment = align_center
            cell_it.border = thin_border
            cell_it.fill = fill_soft_gray
            
            for idx in range(2):
                cell_val = ws.cell(row=r, column=14+idx, value=float(centers_step[idx]))
                cell_val.font = font_regular
                cell_val.number_format = "Rp #,##0"
                cell_val.alignment = align_right
                cell_val.border = thin_border
            
        # Table Header
        headers = [
            "Ranking (No Opsi)", "ID Tempat", "Nama Tempat", "Harga Asli (Rp)", 
            "Jenis Paket", "U_Hemat", "U_Premium", 
            "Target Harga (Rp)", "Selisih ke Target (Rp)", "Rating"
        ]
        
        for c_idx, h in enumerate(headers):
            cell_h = ws.cell(row=11, column=c_idx+1, value=h)
            cell_h.font = font_header
            cell_h.fill = fill_dark
            cell_h.alignment = align_center
            cell_h.border = thin_border
            
        # Populate baris data
        df_sorted = data_pkg["sorted_df"]
        for idx, row in df_sorted.iterrows():
            r = 12 + idx
            
            # Ranking / No Opsi
            cell_rank = ws.cell(row=r, column=1, value=int(row["Ranking"]))
            cell_rank.font = font_bold
            cell_rank.alignment = align_center
            cell_rank.border = thin_border
            cell_rank.fill = fill_light_teal
            
            # ID & Nama
            ws.cell(row=r, column=2, value=int(row["Id_Tempat"])).alignment = align_center
            ws.cell(row=r, column=3, value=row["Nama_Tempat"]).alignment = align_left
            
            # Harga
            cell_pr = ws.cell(row=r, column=4, value=float(row["Estimasi_Harga"]))
            cell_pr.number_format = "Rp #,##0"
            cell_pr.alignment = align_right
            
            # Jenis Paket
            cell_lbl = ws.cell(row=r, column=5, value=str(row["Cluster_Label"]).upper())
            cell_lbl.font = font_bold
            cell_lbl.alignment = align_center
            if row["Cluster_Label"] == "Hemat": cell_lbl.fill = fill_hemat
            else: cell_lbl.fill = fill_premium
            
            # Membership
            ws.cell(row=r, column=6, value=float(row["U_Hemat"])).number_format = "0.0000"
            ws.cell(row=r, column=7, value=float(row["U_Premium"])).number_format = "0.0000"
            
            for c in [6, 7]:
                ws.cell(row=r, column=c).alignment = align_right
            
            # Target Harga & Selisih
            cell_t = ws.cell(row=r, column=8, value=float(row["Target_Price"]) if pd.notna(row["Target_Price"]) else "-")
            if pd.notna(row["Target_Price"]):
                cell_t.number_format = "Rp #,##0"
            cell_t.alignment = align_right
            
            cell_d = ws.cell(row=r, column=9, value=float(row["Selisih_Harga"]) if pd.notna(row["Selisih_Harga"]) else "-")
            if pd.notna(row["Selisih_Harga"]):
                cell_d.number_format = "Rp #,##0"
            cell_d.alignment = align_right
            
            # Rating
            rating_val = row.get("Rating", 0.0)
            cell_rt = ws.cell(row=r, column=10, value=float(rating_val) if pd.notna(rating_val) else 0.0)
            cell_rt.number_format = "0.0"
            cell_rt.alignment = align_center
            
            # Borders & Fonts
            for c in range(1, 11):
                ws.cell(row=r, column=c).border = thin_border
                if c != 1 and c != 5:
                    ws.cell(row=r, column=c).font = font_regular
                    
        auto_fit_columns(ws)
        
    # ==============================================================================
    # SHEET 4: REKOMENDASI PAKET WISATA (TOP 5)
    # ==============================================================================
    ws_pkg = wb.create_sheet("Rekomendasi_Paket")
    ws_pkg.views.sheetView[0].showGridLines = True
    
    # Banner Atas
    ws_pkg.merge_cells("A1:K2")
    ws_pkg["A1"] = f"REKOMENDASI KOMBINASI PAKET WISATA TERPILIH (TOP 5) - {workflow_title.upper()}"
    ws_pkg["A1"].font = font_title
    ws_pkg["A1"].fill = fill_dark
    ws_pkg["A1"].alignment = align_center
    
    # Judul Keterangan Parameter
    ws_pkg["A4"] = f"PARAMETER ACUAN: Anggaran Max = Rp {DEFAULT_BUDGET:,.0f} | Peserta = {DEFAULT_PERSONS} Orang | Durasi = {DEFAULT_DURATION} Hari"
    ws_pkg["A4"].font = font_bold
    ws_pkg.merge_cells("A4:K4")
    
    r_cursor = 6
    
    classes_names = ["PAKET HEMAT (BUDGET)", "PAKET PREMIUM (MEWAH)"]
    
    for idx_cl, cl_name in enumerate(classes_names):
        ws_pkg.cell(row=r_cursor, column=1, value=cl_name).font = font_section
        ws_pkg.merge_cells(start_row=r_cursor, start_column=1, end_row=r_cursor, end_column=11)
        r_cursor += 1
        
        # Header Tabel Paket
        headers_pkg = [
            "No Paket", "Akomodasi (Hotel)", "Tiket Wisata", "Kuliner (Konsumsi)",
            "Total Biaya Hotel (Rp)", "Total Tiket Wisata (Rp)", "Total Kuliner (Rp)",
            "Estimasi Jarak Spasial", "Biaya Transport (Rp)", "TOTAL ANGGARAN PAKET (Rp)", "Status Kelayakan"
        ]
        
        for c_idx, h in enumerate(headers_pkg):
            cell_h = ws_pkg.cell(row=r_cursor, column=c_idx+1, value=h)
            cell_h.font = font_header
            cell_h.fill = fill_dark
            cell_h.alignment = align_center
            cell_h.border = thin_border
            
        r_cursor += 1
        
        # Ambil paket kombinasi dari pkg_res
        combs = pkg_res.get(idx_cl, [])
        
        if not combs:
            ws_pkg.cell(row=r_cursor, column=1, value="Tidak ada kombinasi paket liburan yang memenuhi batas anggaran.").font = font_bold
            ws_pkg.merge_cells(start_row=r_cursor, start_column=1, end_row=r_cursor, end_column=11)
            ws_pkg.cell(row=r_cursor, column=1).alignment = align_center
            ws_pkg.cell(row=r_cursor, column=1).border = thin_border
            r_cursor += 2
            continue
            
        rooms = math.ceil(DEFAULT_PERSONS / 2.0)
        nights = max(DEFAULT_DURATION - 1, 1)
        
        for p_idx, (h_item, w_item, k_item, dist, total_cost) in enumerate(combs):
            r = r_cursor
            
            # No Paket
            cell_no = ws_pkg.cell(row=r, column=1, value=f"Opsi Paket {p_idx+1}")
            cell_no.font = font_bold
            cell_no.alignment = align_center
            cell_no.border = thin_border
            cell_no.fill = fill_light_teal
            
            # Item Names
            ws_pkg.cell(row=r, column=2, value=h_item["Nama_Tempat"]).alignment = align_left
            ws_pkg.cell(row=r, column=3, value=w_item["Nama_Tempat"]).alignment = align_left
            ws_pkg.cell(row=r, column=4, value=k_item["Nama_Tempat"]).alignment = align_left
            
            # Breakdown Biaya (Hotel, Wisata, Kuliner)
            c_hotel = h_item["Estimasi_Harga"] * nights * rooms
            cell_ch = ws_pkg.cell(row=r, column=5, value=float(c_hotel))
            cell_ch.number_format = "Rp #,##0"
            cell_ch.alignment = align_right
            
            c_wis = w_item["Estimasi_Harga"] * DEFAULT_PERSONS
            cell_cw = ws_pkg.cell(row=r, column=6, value=float(c_wis))
            cell_cw.number_format = "Rp #,##0"
            cell_cw.alignment = align_right
            
            c_kul = k_item["Estimasi_Harga"] * DEFAULT_PERSONS * 3 * DEFAULT_DURATION
            cell_ck = ws_pkg.cell(row=r, column=7, value=float(c_kul))
            cell_ck.number_format = "Rp #,##0"
            cell_ck.alignment = align_right
            
            # Jarak Spasial
            cell_dst = ws_pkg.cell(row=r, column=8, value=f"{dist:.2f} km")
            cell_dst.alignment = align_center
            
            # Transport
            c_trans, _ = get_transport_cost(DEFAULT_PERSONS, dist)
            cell_ctr = ws_pkg.cell(row=r, column=9, value=float(c_trans))
            cell_ctr.number_format = "Rp #,##0"
            cell_ctr.alignment = align_right
            
            # Total Cost
            cell_tot = ws_pkg.cell(row=r, column=10, value=float(total_cost))
            cell_tot.font = font_bold
            cell_tot.number_format = "Rp #,##0"
            cell_tot.alignment = align_right
            cell_tot.fill = fill_light_blue
            
            # Kelayakan (Dinamis / Riel)
            status_str = "LAYAK (OK)" if total_cost <= DEFAULT_BUDGET else "LIMIT ANGGARAN"
            cell_st = ws_pkg.cell(row=r, column=11, value=status_str)
            cell_st.font = font_bold
            cell_st.alignment = align_center
            if status_str.startswith("LAYAK"):
                cell_st.fill = fill_hemat
            else:
                cell_st.fill = fill_premium
                
            # Borders & Fonts
            for c in range(1, 12):
                ws_pkg.cell(row=r, column=c).border = thin_border
                if c != 1 and c != 10 and c != 11:
                    ws_pkg.cell(row=r, column=c).font = font_regular
                    
            r_cursor += 1
            
        r_cursor += 2
        
    auto_fit_columns(ws_pkg)
        
    filepath = os.path.join(output_dir, filename)
    wb.save(filepath)
    print(f"💾 File berhasil disimpan: {filepath}")

# ==============================================================================
# 8. PROCESSING WORKFLOWS & DATA PREPARATION
# ==============================================================================

# ------------------------------------------------------------------------------
# WORKFLOW 1: BUDGET-FIRST
# ------------------------------------------------------------------------------
print("⚡ Memproses data untuk [Workflow 1: Budget-First]...")

def prepare_budget_first_data(df, prices, cntr, U, lbl, anchor, iters_val, centers_history):
    ratios = RATIO_SCHEMES[DEFAULT_SCHEME]
    xb, sigma, sep = calculate_xie_beni_metrics(prices, cntr, U)
    
    # Buat Dataframe hasil
    df_res = df.copy()
    df_res["Cluster"] = lbl
    df_res["Cluster_Label"] = [CLUSTER_LABELS[l] for l in lbl]
    df_res["U_Hemat"] = U[0, :]
    df_res["U_Premium"] = U[1, :]
    
    # Target prices
    targets = [anchor * r for r in ratios]
    df_res["Target_Price"] = [targets[l] for l in lbl]
    df_res["Selisih_Harga"] = (df_res["Estimasi_Harga"] - df_res["Target_Price"]).abs()
    
    # Perangkingan: Diurutkan berdasarkan Cluster (Hemat -> Premium),
    # dan di dalam setiap klaster, diurutkan berdasarkan Selisih Harga ke Target (Terdekat)
    df_res = df_res.sort_values(by=["Cluster", "Selisih_Harga"], ascending=[True, True]).reset_index(drop=True)
    
    # Tambah kolom Ranking (No Opsi)
    rankings = []
    for c_idx in [0, 1]:
        sub_df = df_res[df_res["Cluster"] == c_idx]
        rankings.extend(range(1, len(sub_df) + 1))
    df_res["Ranking"] = rankings
    
    return {
        "n": len(df),
        "iters": iters_val,
        "sigma": sigma,
        "sep": sep,
        "xb": xb,
        "centroids": cntr,
        "anchors": targets,
        "centers_history": centers_history,
        "sorted_df": df_res
    }

def prepare_flexible_data(df, prices, cntr, U, lbl, iters_val, centers_history):
    xb, sigma, sep = calculate_xie_beni_metrics(prices, cntr, U)
    
    df_res = df.copy()
    df_res["Cluster"] = lbl
    df_res["Cluster_Label"] = [CLUSTER_LABELS[l] for l in lbl]
    df_res["U_Hemat"] = U[0, :]
    df_res["U_Premium"] = U[1, :]
    
    # Target price di flexible adalah centroid itu sendiri
    df_res["Target_Price"] = [cntr[l] for l in lbl]
    df_res["Selisih_Harga"] = (df_res["Estimasi_Harga"] - df_res["Target_Price"]).abs()
    
    # Untuk ranking di flexible, diurutkan berdasarkan Cluster, dan di dalam setiap klaster,
    # diurutkan berdasarkan Derajat Keanggotaan tertinggi ke terendah (Membership Degree)
    membership_degrees = []
    for idx, row in df_res.iterrows():
        c_idx = row["Cluster"]
        membership_degrees.append(U[c_idx, idx])
    df_res["Membership_Degree"] = membership_degrees
    
    df_res = df_res.sort_values(by=["Cluster", "Membership_Degree"], ascending=[True, False]).reset_index(drop=True)
    
    rankings = []
    for c_idx in [0, 1]:
        sub_df = df_res[df_res["Cluster"] == c_idx]
        rankings.extend(range(1, len(sub_df) + 1))
    df_res["Ranking"] = rankings
    
    return {
        "n": len(df),
        "iters": iters_val,
        "sigma": sigma,
        "sep": sep,
        "xb": xb,
        "centroids": cntr,
        "anchors": None,
        "centers_history": centers_history,
        "sorted_df": df_res
    }

def prepare_dest_first_wisata_data(df, prices, cntr, U, lbl, locked_item, iters_val, centers_history):
    xb, sigma, sep = calculate_xie_beni_metrics(prices, cntr, U)
    
    df_res = df.copy()
    df_res["Cluster"] = lbl
    df_res["Cluster_Label"] = [CLUSTER_LABELS[l] for l in lbl]
    df_res["U_Hemat"] = U[0, :]
    df_res["U_Premium"] = U[1, :]
    
    # Target price & selisih
    df_res["Target_Price"] = np.nan
    df_res["Selisih_Harga"] = np.nan
    
    # Hitung Jarak Haversine spasial ke locked_item
    lat_target = locked_item["Latitude"]
    lon_target = locked_item["Longitude"]
    
    distances = []
    for idx, row in df_res.iterrows():
        d = haversine_distance(row["Latitude"], row["Longitude"], lat_target, lon_target)
        distances.append(d)
    df_res["Distance_To_Locked"] = distances
    
    # Diurutkan berdasarkan Jarak terdekat
    df_res = df_res.sort_values(by="Distance_To_Locked", ascending=True).reset_index(drop=True)
    df_res["Ranking"] = range(1, len(df_res) + 1)
    
    return {
        "n": len(df),
        "iters": iters_val,
        "sigma": sigma,
        "sep": sep,
        "xb": xb,
        "centroids": cntr,
        "anchors": None,
        "centers_history": centers_history,
        "sorted_df": df_res
    }

w_bf_data_temp = prepare_budget_first_data(df_wisata, df_wisata["Estimasi_Harga"].values, np.zeros((2,1)), np.zeros((2, len(df_wisata))), np.zeros(len(df_wisata), dtype=int), 0, 0, [])

# ==============================================================================
# 9. RUNNING THE 7 STRESS-TESTING SCENARIOS SEQUENTIALLY
# ==============================================================================
print("\n🚀 Memulai Eksekusi 7 Skenario Pengujian Batas Ekstrem (Stress-Testing)...")

import shutil
if os.path.exists(output_dir):
    shutil.rmtree(output_dir)
os.makedirs(output_dir)
print(f"📁 Folder '{output_dir}' dibersihkan dan disiapkan.")

# ------------------------------------------------------------------------------
# SCENARIO 1: WORKFLOW 1 (BUDGET-FIRST)
# ------------------------------------------------------------------------------

# 1a. Budget Termurah: Rp 27.750
print("\n⚡ Menjalankan Skenario [1a] Budget-First Termurah (Rp 27.750)...")
DEFAULT_BUDGET = 27750.0

alloc_akomodasi = DEFAULT_BUDGET * 0.40
alloc_wisata = DEFAULT_BUDGET * 0.15
alloc_kuliner = DEFAULT_BUDGET * 0.20

rooms = math.ceil(DEFAULT_PERSONS / 2.0)
nights = max(DEFAULT_DURATION - 1, 1)

anchor_hotel = alloc_akomodasi / (nights * rooms)
anchor_wisata = alloc_wisata / DEFAULT_PERSONS
anchor_kuliner = alloc_kuliner / (DEFAULT_PERSONS * 3 * DEFAULT_DURATION)

# Run FCMs
w_cntr_bgt, w_u_bgt, w_lbl_bgt, w_it_bgt, w_his_bgt = run_budget_anchored_fcm(df_wisata["Estimasi_Harga"].values, anchor_wisata, DEFAULT_SCHEME)
h_cntr_bgt, h_u_bgt, h_lbl_bgt, h_it_bgt, h_his_bgt = run_budget_anchored_fcm(df_hotel["Estimasi_Harga"].values, anchor_hotel, DEFAULT_SCHEME)
k_cntr_bgt, k_u_bgt, k_lbl_bgt, k_it_bgt, k_his_bgt = run_budget_anchored_fcm(df_kuliner["Estimasi_Harga"].values, anchor_kuliner, DEFAULT_SCHEME)

w_bf_data = prepare_budget_first_data(df_wisata, df_wisata["Estimasi_Harga"].values, w_cntr_bgt, w_u_bgt, w_lbl_bgt, anchor_wisata, w_it_bgt, w_his_bgt)
h_bf_data = prepare_budget_first_data(df_hotel, df_hotel["Estimasi_Harga"].values, h_cntr_bgt, h_u_bgt, h_lbl_bgt, anchor_hotel, h_it_bgt, h_his_bgt)
k_bf_data = prepare_budget_first_data(df_kuliner, df_kuliner["Estimasi_Harga"].values, k_cntr_bgt, k_u_bgt, k_lbl_bgt, anchor_kuliner, k_it_bgt, k_his_bgt)

pkg_bf_min = get_budget_first_packages(df_wisata, df_hotel, df_kuliner, w_lbl_bgt, h_lbl_bgt, k_lbl_bgt, w_cntr_bgt, h_cntr_bgt, k_cntr_bgt, anchor_wisata, anchor_hotel, anchor_kuliner, RATIO_SCHEMES[DEFAULT_SCHEME])
generate_ranking_excel("1a_BudgetFirst_Termurah_27k.xlsx", "Budget-First (Rp 27.750)", w_bf_data, h_bf_data, k_bf_data, pkg_bf_min)


# 1b. Budget Termahal: Rp 5.770.750
print("⚡ Menjalankan Skenario [1b] Budget-First Termahal (Rp 5.770.750)...")
DEFAULT_BUDGET = 5770750.0

alloc_akomodasi = DEFAULT_BUDGET * 0.40
alloc_wisata = DEFAULT_BUDGET * 0.15
alloc_kuliner = DEFAULT_BUDGET * 0.20

anchor_hotel = alloc_akomodasi / (nights * rooms)
anchor_wisata = alloc_wisata / DEFAULT_PERSONS
anchor_kuliner = alloc_kuliner / (DEFAULT_PERSONS * 3 * DEFAULT_DURATION)

w_cntr_bgt, w_u_bgt, w_lbl_bgt, w_it_bgt, w_his_bgt = run_budget_anchored_fcm(df_wisata["Estimasi_Harga"].values, anchor_wisata, DEFAULT_SCHEME)
h_cntr_bgt, h_u_bgt, h_lbl_bgt, h_it_bgt, h_his_bgt = run_budget_anchored_fcm(df_hotel["Estimasi_Harga"].values, anchor_hotel, DEFAULT_SCHEME)
k_cntr_bgt, k_u_bgt, k_lbl_bgt, k_it_bgt, k_his_bgt = run_budget_anchored_fcm(df_kuliner["Estimasi_Harga"].values, anchor_kuliner, DEFAULT_SCHEME)

w_bf_data = prepare_budget_first_data(df_wisata, df_wisata["Estimasi_Harga"].values, w_cntr_bgt, w_u_bgt, w_lbl_bgt, anchor_wisata, w_it_bgt, w_his_bgt)
h_bf_data = prepare_budget_first_data(df_hotel, df_hotel["Estimasi_Harga"].values, h_cntr_bgt, h_u_bgt, h_lbl_bgt, anchor_hotel, h_it_bgt, h_his_bgt)
k_bf_data = prepare_budget_first_data(df_kuliner, df_kuliner["Estimasi_Harga"].values, k_cntr_bgt, k_u_bgt, k_lbl_bgt, anchor_kuliner, k_it_bgt, k_his_bgt)

pkg_bf_max = get_budget_first_packages(df_wisata, df_hotel, df_kuliner, w_lbl_bgt, h_lbl_bgt, k_lbl_bgt, w_cntr_bgt, h_cntr_bgt, k_cntr_bgt, anchor_wisata, anchor_hotel, anchor_kuliner, RATIO_SCHEMES[DEFAULT_SCHEME])
generate_ranking_excel("1b_BudgetFirst_Termahal_5M.xlsx", "Budget-First (Rp 5.770.750)", w_bf_data, h_bf_data, k_bf_data, pkg_bf_max)


# ------------------------------------------------------------------------------
# SCENARIO 2: WORKFLOW 2 (FLEXIBLE EXPLORATION)
# ------------------------------------------------------------------------------
print("\n⚡ Menjalankan Skenario [2] Flexible Exploration...")
DEFAULT_BUDGET = 1500000.0  # Reset ke nominal default sebagai patokan paket

w_cntr_std, w_u_std, w_lbl_std, w_it_std, w_his_std = run_percentile_fcm(df_wisata["Estimasi_Harga"].values)
h_cntr_std, h_u_std, h_lbl_std, h_it_std, h_his_std = run_percentile_fcm(df_hotel["Estimasi_Harga"].values)
k_cntr_std, k_u_std, k_lbl_std, k_it_std, k_his_std = run_percentile_fcm(df_kuliner["Estimasi_Harga"].values)

w_fl_data = prepare_flexible_data(df_wisata, df_wisata["Estimasi_Harga"].values, w_cntr_std, w_u_std, w_lbl_std, w_it_std, w_his_std)
h_fl_data = prepare_flexible_data(df_hotel, df_hotel["Estimasi_Harga"].values, h_cntr_std, h_u_std, h_lbl_std, h_it_std, h_his_std)
k_fl_data = prepare_flexible_data(df_kuliner, df_kuliner["Estimasi_Harga"].values, k_cntr_std, k_u_std, k_lbl_std, k_it_std, k_his_std)

pkg_fl = get_flexible_packages(df_wisata, df_hotel, df_kuliner, w_lbl_std, h_lbl_std, k_lbl_std, w_cntr_std, h_cntr_std, k_cntr_std)
generate_ranking_excel("2_FlexibleExploration.xlsx", "Flexible Exploration", w_fl_data, h_fl_data, k_fl_data, pkg_fl)


# ------------------------------------------------------------------------------
# SCENARIO 3: WORKFLOW 3 (DESTINATION-FIRST)
# ------------------------------------------------------------------------------

# 3a. Kondisi A: Spasial Ekstrem Tanpa Input Budget (Default Rp 1.500.000)

# 3a1. Tempat Termurah: ID 40 Gua Maria Sendang Purwaningsih (HTM Rp 0)
print("\n⚡ Menjalankan Skenario [3a1] Destination-First Termurah (HTM Rp 0, Budget Default Rp 1.500.000)...")
DEFAULT_BUDGET = 1500000.0
locked_dest_min = df_wisata[df_wisata["Id_Tempat"] == CHEAPEST_WISATA_ID].iloc[0].to_dict()

# Hitung ulang alokasi sisa budget setelah dikurangi tiket wisata utama (Rp 0)
sisa_budget = DEFAULT_BUDGET - (locked_dest_min["Estimasi_Harga"] * DEFAULT_PERSONS)
anchor_hotel_dest = (sisa_budget * (40.0/85.0)) / (nights * rooms)
anchor_kuliner_dest = (sisa_budget * (20.0/85.0)) / (DEFAULT_PERSONS * 3 * DEFAULT_DURATION)

# FCM khusus
h_cntr_dest, h_u_dest, h_lbl_dest, h_it_dest, h_his_dest = run_budget_anchored_fcm(df_hotel["Estimasi_Harga"].values, anchor_hotel_dest, DEFAULT_SCHEME)
k_cntr_dest, k_u_dest, k_lbl_dest, k_it_dest, k_his_dest = run_budget_anchored_fcm(df_kuliner["Estimasi_Harga"].values, anchor_kuliner_dest, DEFAULT_SCHEME)

w_ds_data = prepare_dest_first_wisata_data(df_wisata, df_wisata["Estimasi_Harga"].values, w_cntr_bgt, w_u_bgt, w_lbl_bgt, locked_dest_min, w_it_bgt, w_his_bgt)
h_ds_data = prepare_budget_first_data(df_hotel, df_hotel["Estimasi_Harga"].values, h_cntr_dest, h_u_dest, h_lbl_dest, anchor_hotel_dest, h_it_dest, h_his_dest)
k_ds_data = prepare_budget_first_data(df_kuliner, df_kuliner["Estimasi_Harga"].values, k_cntr_dest, k_u_dest, k_lbl_dest, anchor_kuliner_dest, k_it_dest, k_his_dest)

pkg_ds_3a1 = get_dest_first_packages(df_wisata, df_hotel, df_kuliner, h_lbl_dest, k_lbl_dest, h_cntr_dest, k_cntr_dest, locked_dest_min, anchor_hotel_dest, anchor_kuliner_dest, RATIO_SCHEMES[DEFAULT_SCHEME])
generate_ranking_excel("3a1_DestinationFirst_Termurah_NoBudget.xlsx", "Destination-First (ID 40 Termurah, Rp 1.500.000)", w_ds_data, h_ds_data, k_ds_data, pkg_ds_3a1)


# 3a2. Tempat Termahal: ID 159 Pujon Adventure & Rafting (HTM Rp 275.000)
print("⚡ Menjalankan Skenario [3a2] Destination-First Termahal (HTM Rp 275.000, Budget Default Rp 1.500.000)...")
locked_dest_max = df_wisata[df_wisata["Id_Tempat"] == EXPENSIVE_WISATA_ID].iloc[0].to_dict()

sisa_budget = DEFAULT_BUDGET - (locked_dest_max["Estimasi_Harga"] * DEFAULT_PERSONS)
anchor_hotel_dest = (sisa_budget * (40.0/85.0)) / (nights * rooms)
anchor_kuliner_dest = (sisa_budget * (20.0/85.0)) / (DEFAULT_PERSONS * 3 * DEFAULT_DURATION)

h_cntr_dest, h_u_dest, h_lbl_dest, h_it_dest, h_his_dest = run_budget_anchored_fcm(df_hotel["Estimasi_Harga"].values, anchor_hotel_dest, DEFAULT_SCHEME)
k_cntr_dest, k_u_dest, k_lbl_dest, k_it_dest, k_his_dest = run_budget_anchored_fcm(df_kuliner["Estimasi_Harga"].values, anchor_kuliner_dest, DEFAULT_SCHEME)

w_ds_data = prepare_dest_first_wisata_data(df_wisata, df_wisata["Estimasi_Harga"].values, w_cntr_bgt, w_u_bgt, w_lbl_bgt, locked_dest_max, w_it_bgt, w_his_bgt)
h_ds_data = prepare_budget_first_data(df_hotel, df_hotel["Estimasi_Harga"].values, h_cntr_dest, h_u_dest, h_lbl_dest, anchor_hotel_dest, h_it_dest, h_his_dest)
k_ds_data = prepare_budget_first_data(df_kuliner, df_kuliner["Estimasi_Harga"].values, k_cntr_dest, k_u_dest, k_lbl_dest, anchor_kuliner_dest, k_it_dest, k_his_dest)

pkg_ds_3a2 = get_dest_first_packages(df_wisata, df_hotel, df_kuliner, h_lbl_dest, k_lbl_dest, h_cntr_dest, k_cntr_dest, locked_dest_max, anchor_hotel_dest, anchor_kuliner_dest, RATIO_SCHEMES[DEFAULT_SCHEME])
generate_ranking_excel("3a2_DestinationFirst_Termahal_NoBudget.xlsx", "Destination-First (ID 159 Termahal, Rp 1.500.000)", w_ds_data, h_ds_data, k_ds_data, pkg_ds_3a2)


# 3b. Kondisi B: Spasial Ekstrem Dengan Input Budget Ekstrem

# 3b1. Tempat Termurah + Budget Termurah Rp 27.750
print("\n⚡ Menjalankan Skenario [3b1] Destination-First Termurah (HTM Rp 0, Budget Termurah Rp 27.750)...")
DEFAULT_BUDGET = 27750.0

sisa_budget = DEFAULT_BUDGET - (locked_dest_min["Estimasi_Harga"] * DEFAULT_PERSONS)
anchor_hotel_dest = (sisa_budget * (40.0/85.0)) / (nights * rooms)
anchor_kuliner_dest = (sisa_budget * (20.0/85.0)) / (DEFAULT_PERSONS * 3 * DEFAULT_DURATION)

h_cntr_dest, h_u_dest, h_lbl_dest, h_it_dest, h_his_dest = run_budget_anchored_fcm(df_hotel["Estimasi_Harga"].values, anchor_hotel_dest, DEFAULT_SCHEME)
k_cntr_dest, k_u_dest, k_lbl_dest, k_it_dest, k_his_dest = run_budget_anchored_fcm(df_kuliner["Estimasi_Harga"].values, anchor_kuliner_dest, DEFAULT_SCHEME)

w_ds_data = prepare_dest_first_wisata_data(df_wisata, df_wisata["Estimasi_Harga"].values, w_cntr_bgt, w_u_bgt, w_lbl_bgt, locked_dest_min, w_it_bgt, w_his_bgt)
h_ds_data = prepare_budget_first_data(df_hotel, df_hotel["Estimasi_Harga"].values, h_cntr_dest, h_u_dest, h_lbl_dest, anchor_hotel_dest, h_it_dest, h_his_dest)
k_ds_data = prepare_budget_first_data(df_kuliner, df_kuliner["Estimasi_Harga"].values, k_cntr_dest, k_u_dest, k_lbl_dest, anchor_kuliner_dest, k_it_dest, k_his_dest)

pkg_ds_3b1 = get_dest_first_packages(df_wisata, df_hotel, df_kuliner, h_lbl_dest, k_lbl_dest, h_cntr_dest, k_cntr_dest, locked_dest_min, anchor_hotel_dest, anchor_kuliner_dest, RATIO_SCHEMES[DEFAULT_SCHEME])
generate_ranking_excel("3b1_DestinationFirst_Termurah_Budget27k.xlsx", "Destination-First (ID 40 Termurah, Rp 27.750)", w_ds_data, h_ds_data, k_ds_data, pkg_ds_3b1)


# 3b2. Tempat Termahal + Budget Termahal Rp 5.770.750
print("⚡ Menjalankan Skenario [3b2] Destination-First Termahal (HTM Rp 275.000, Budget Termahal Rp 5.770.750)...")
DEFAULT_BUDGET = 5770750.0

sisa_budget = DEFAULT_BUDGET - (locked_dest_max["Estimasi_Harga"] * DEFAULT_PERSONS)
anchor_hotel_dest = (sisa_budget * (40.0/85.0)) / (nights * rooms)
anchor_kuliner_dest = (sisa_budget * (20.0/85.0)) / (DEFAULT_PERSONS * 3 * DEFAULT_DURATION)

h_cntr_dest, h_u_dest, h_lbl_dest, h_it_dest, h_his_dest = run_budget_anchored_fcm(df_hotel["Estimasi_Harga"].values, anchor_hotel_dest, DEFAULT_SCHEME)
k_cntr_dest, k_u_dest, k_lbl_dest, k_it_dest, k_his_dest = run_budget_anchored_fcm(df_kuliner["Estimasi_Harga"].values, anchor_kuliner_dest, DEFAULT_SCHEME)

w_ds_data = prepare_dest_first_wisata_data(df_wisata, df_wisata["Estimasi_Harga"].values, w_cntr_bgt, w_u_bgt, w_lbl_bgt, locked_dest_max, w_it_bgt, w_his_bgt)
h_ds_data = prepare_budget_first_data(df_hotel, df_hotel["Estimasi_Harga"].values, h_cntr_dest, h_u_dest, h_lbl_dest, anchor_hotel_dest, h_it_dest, h_his_dest)
k_ds_data = prepare_budget_first_data(df_kuliner, df_kuliner["Estimasi_Harga"].values, k_cntr_dest, k_u_dest, k_lbl_dest, anchor_kuliner_dest, k_it_dest, k_his_dest)

pkg_ds_3b2 = get_dest_first_packages(df_wisata, df_hotel, df_kuliner, h_lbl_dest, k_lbl_dest, h_cntr_dest, k_cntr_dest, locked_dest_max, anchor_hotel_dest, anchor_kuliner_dest, RATIO_SCHEMES[DEFAULT_SCHEME])
generate_ranking_excel("3b2_DestinationFirst_Termahal_Budget5M.xlsx", "Destination-First (ID 159 Termahal, Rp 5.770.750)", w_ds_data, h_ds_data, k_ds_data, pkg_ds_3b2)


print(f"\n🎉 SUKSES 100%! SEBANYAK 7 FILE EXCEL HASIL RANKING PENGUJIAN BATAS EKSTREM TELAH BERHASIL DICETAK DI FOLDER '{output_dir}'!")
